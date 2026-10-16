#!/usr/bin/env python3
"""Make the empty office workbook "Hot Lunch Orders.xlsx" (python3 worker/make_workbook.py).

It goes in the office mailbox's OneDrive (root folder). The order service adds one row per
student to the "Orders" table; its columns must stay in this order.
"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = pathlib.Path(__file__).with_name("Hot Lunch Orders.xlsx")
COLUMNS = ["Submitted", "Week of", "Order #", "Parent", "Email", "Phone", "Student", "Grade", "Preschool",
           "Price per lunch", "Mon", "Tue", "Wed", "Thu", "Fri", "Lunches", "Total", "Notes"]
WIDTHS = [18, 11, 9, 20, 26, 14, 22, 16, 10, 9, 6, 6, 6, 6, 6, 9, 10, 30]
NAVY, GOLD = "1B2A4A", "C9972C"

wb = Workbook()
ws = wb.active
ws.title = "Orders"
ws.append(COLUMNS)
# A table needs one data row; the order service adds the real ones below it.
ws.append([None] * 15 + ["=COUNTIF(Orders[[#This Row],[Mon]:[Fri]],\"X\")",
                         "=Orders[[#This Row],[Lunches]]*Orders[[#This Row],[Price per lunch]]", None])
table = Table(displayName="Orders", ref=f"A1:R2")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(table)
for i, w in enumerate(WIDTHS, start=1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
ws["B2"].number_format = "m/d/yyyy"
for c in ("J", "Q"):
    for r in (1, 2):
        ws[f"{c}{r}"].number_format = '"$"#,##0.00'
ws.freeze_panes = "A2"

k = wb.create_sheet("Kitchen")
k["A1"], k["A1"].font = "Hot lunch counts", Font(bold=True, size=14, color=NAVY)
k["A3"], k["B3"] = "Week of (Monday)", "=MAX(Orders[Week of])"
k["A3"].font = Font(bold=True)
k["B3"].number_format = "m/d/yyyy"
k["B3"].fill = PatternFill("solid", fgColor="FFF7E0")
k["C3"] = "Shows the latest week. Type another Monday's date here to see that week."
k["C3"].font = Font(italic=True, color="666666")
k.append([])
k.append(["Day", "Date", "K5–12th", "Preschool", "Total"])
for c in "ABCDE":
    k[f"{c}5"].font = Font(bold=True, color="FFFFFF")
    k[f"{c}5"].fill = PatternFill("solid", fgColor=NAVY)
for i, (day, col) in enumerate([("Monday", "Mon"), ("Tuesday", "Tue"), ("Wednesday", "Wed"),
                                ("Thursday", "Thu"), ("Friday", "Fri")]):
    r = 6 + i
    k[f"A{r}"] = day
    k[f"B{r}"] = f"=$B$3+{i}"
    k[f"B{r}"].number_format = "ddd m/d"
    k[f"C{r}"] = f'=COUNTIFS(Orders[Week of],$B$3,Orders[{col}],"X",Orders[Preschool],"No")'
    k[f"D{r}"] = f'=COUNTIFS(Orders[Week of],$B$3,Orders[{col}],"X",Orders[Preschool],"Yes")'
    k[f"E{r}"] = f"=C{r}+D{r}"
k["A11"], k["E11"] = "Week total", "=SUM(E6:E10)"
k["A11"].font = k["E11"].font = Font(bold=True)
k["A12"], k["E12"] = "Week charges", '=SUMIFS(Orders[Total],Orders[Week of],$B$3)'
k["E12"].number_format = '"$"#,##0.00'
for col, w in zip("ABCDE", (18, 12, 10, 11, 10)):
    k.column_dimensions[col].width = w

h = wb.create_sheet("How to use")
h.column_dimensions["A"].width = 110
lines = [
    ("Hot lunch orders", True),
    ("Orders placed on the school website land on the Orders tab by themselves: one row per student, "
     "with an X under each day ordered. Lunches and Total add themselves up.", False),
    ("", False),
    ("Kitchen counts: open the Kitchen tab. It shows the latest week; type another Monday's date in the "
     "yellow box to see a different week.", False),
    ("Billing: on the Orders tab, use the filter on 'Week of' to show one week. Each row's Total is what "
     "to charge that student in Gradelink.", False),
    ("Changes by phone: remove an X (or type X) under a day and the totals update. To cancel a whole order, "
     "delete the student's row. Use Notes for anything worth remembering.", False),
    ("", False),
    ("Please don't rename this file, the Orders tab, the Orders table or its column headings, and don't move "
     "the file out of the OneDrive's main folder. The website needs them to add orders.", True),
]
for i, (text, bold) in enumerate(lines, start=1):
    h[f"A{i}"] = text
    h[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    h[f"A{i}"].font = Font(bold=bold, size=14 if i == 1 else 11, color=NAVY if i == 1 else "000000")

wb.save(OUT)
print("wrote", OUT)
